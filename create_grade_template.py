"""生成空白成绩导入模板；不读取学生名单或包含示例学生。"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from gradebook import SUBJECTS


def create(filename):
    wb=Workbook(); ws=wb.active; ws.title="成绩"
    columns=["学生姓名"]
    for subject in SUBJECTS: columns.extend([subject+"成绩", "班名次", "校名次"])
    columns.extend(["总分", "班名次", "校名次", "年级人数"])
    ws.append(columns); ws.freeze_panes="B2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(columns))}1"
    for i in range(1,len(columns)+1):
        cell=ws.cell(1,i); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="356B8C")
        ws.column_dimensions[get_column_letter(i)].width=18
    note=wb.create_sheet("填表说明")
    for text in ["按学生姓名匹配，学号不作为成绩的唯一键。姓名重复必须先由老师区分。",
                 "一行一名学生，成绩表只保留单行表头；姓名也可以使用“姓名”列名。",
                 "科目/科目成绩后面的班名次、班次、班排名都属于该科；校名次、校排名、校次同理。遇到下一科分数列后切换归属。",
                 "例如：物理成绩 | 班次 | 校次 | 化学 | 班排名 | 校排名。也支持完整列名“物理班名次”“物理校排名”。",
                 "校排名填写真实学校/年级名次，不是班级排名。缺少校排名可留空，系统不编造。",
                 "缺考、未考或空白不计为 0 分；真实 0 分请填写数字 0。",
                 "建议填写总分及总分校排名。不会自动把部分科目分数求和当作总分。",
                 "可选年级人数用于提醒参考人数变化；校名次不得超过该人数。",
                 "考试名称确认后按导入时间排序，不按学号、文件创建时间或文件名排序。",
                 "若使用公式，请先在 Excel 中重新计算并保存，系统只读取缓存的计算结果。",
                 "导入后先预览确认，不会自动新增未匹配学生或立即换座。"]:
        note.append([text])
    note.column_dimensions["A"].width=110
    wb.save(filename); wb.close()


if __name__=="__main__":create(Path(__file__).with_name("成绩导入模板.xlsx"))
