#!/usr/bin/env python3
"""离线学生信息与周末住宿管理系统（命令行版）。"""
import csv, json, os, sys
from datetime import datetime

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = load_workbook = None

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE, "students.json")
FIELDS = ["学号", "姓名", "班级", "性别", "联系电话", "家长电话", "周末住宿", "住宿地点", "备注"]

def load_data():
    if not os.path.exists(DATA_FILE): return []
    with open(DATA_FILE, encoding="utf-8") as f: return json.load(f)

def save_data(rows):
    tmp = DATA_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f: json.dump(rows, f, ensure_ascii=False, indent=2)
    os.replace(tmp, DATA_FILE)

def yn(value): return str(value).strip().lower() in ("是", "y", "yes", "1", "true")
def input_student(old=None):
    old = old or {}
    def ask(k, label, required=False):
        while True:
            v = input(f"{label}[{old.get(k,'')}]: ").strip() or str(old.get(k, ""))
            if not required or v: return v
            print("此项不能为空")
    return {"学号": ask("学号", "学号", True), "姓名": ask("姓名", "姓名", True),
            "班级": ask("班级", "班级"), "性别": ask("性别", "性别"),
            "联系电话": ask("联系电话", "联系电话"), "家长电话": ask("家长电话", "家长电话"),
            "周末住宿": "是" if yn(ask("周末住宿", "周末住宿(是/否)")) else "否",
            "住宿地点": ask("住宿地点", "住宿地点"), "备注": ask("备注", "备注")}

def add(rows):
    s=input_student()
    if any(x["学号"]==s["学号"] for x in rows): print("学号已存在"); return
    rows.append(s); save_data(rows); print("添加成功")
def edit(rows):
    sid=input("请输入要修改的学号：").strip(); i=next((i for i,x in enumerate(rows) if x["学号"]==sid),None)
    if i is None: print("未找到"); return
    rows[i]=input_student(rows[i]); save_data(rows); print("修改成功")
def remove(rows):
    sid=input("请输入要删除的学号：").strip(); n=len(rows); rows[:]=[x for x in rows if x["学号"]!=sid]
    if len(rows)<n: save_data(rows); print("删除成功")
    else: print("未找到")
def list_rows(rows, only_weekend=False):
    data=[x for x in rows if not only_weekend or x.get("周末住宿")=="是"]
    print(f"共 {len(data)} 人")
    for x in data: print(" | ".join(str(x.get(k,"")) for k in FIELDS))
def import_file(rows):
    path=input("CSV/XLSX文件路径：").strip().strip('"')
    try:
        if path.lower().endswith(".xlsx"):
            if not load_workbook: raise RuntimeError("请安装 openpyxl：pip install openpyxl")
            ws=load_workbook(path, read_only=True, data_only=True).active; vals=list(ws.values); headers=list(vals[0]); records=[dict(zip(headers,r)) for r in vals[1:]]
        else:
            with open(path,encoding="utf-8-sig",newline="") as f: records=list(csv.DictReader(f))
        existing={x["学号"]:x for x in rows}; count=0
        for r in records:
            sid=str(r.get("学号","")).strip()
            if not sid: continue
            s={k:str(r.get(k,"") or "").strip() for k in FIELDS}; s["学号"]=sid; s["周末住宿"]="是" if yn(s["周末住宿"]) else "否"; existing[sid]=s; count+=1
        rows[:]=list(existing.values()); save_data(rows); print(f"导入完成，处理 {count} 条，当前共 {len(rows)} 人")
    except Exception as e: print(f"导入失败：{e}")
def export_xlsx(rows):
    if not Workbook: print("请安装 openpyxl：pip install openpyxl"); return
    path=input("导出文件路径（默认 weekend_report.xlsx）：").strip() or os.path.join(BASE,"weekend_report.xlsx")
    wb=Workbook(); ws=wb.active; ws.title="学生信息"; ws.append(FIELDS)
    for x in rows: ws.append([x.get(k,"") for k in FIELDS])
    ws2=wb.create_sheet("周末住宿名单"); ws2.append(FIELDS)
    for x in rows:
        if x.get("周末住宿")=="是": ws2.append([x.get(k,"") for k in FIELDS])
    wb.save(path); print(f"已导出：{path}")
def main():
    rows=load_data(); print("学生信息与周末住宿管理系统（离线版）")
    actions={"1":lambda:add(rows),"2":lambda:edit(rows),"3":lambda:remove(rows),"4":lambda:list_rows(rows),"5":lambda:list_rows(rows,True),"6":lambda:import_file(rows),"7":lambda:export_xlsx(rows)}
    while True:
        print("\n1添加 2修改 3删除 4查看全部 5周末住宿统计/名单 6导入CSV/XLSX 7导出XLSX 0退出")
        c=input("请选择：").strip()
        if c=="0": break
        if c in actions:
            try: actions[c]()
            except (EOFError,KeyboardInterrupt): print(); break
        else: print("无效选项")
if __name__=="__main__": main()
